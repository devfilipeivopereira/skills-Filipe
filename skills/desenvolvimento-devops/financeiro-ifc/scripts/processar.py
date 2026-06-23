#!/usr/bin/env python3
"""Trata planilhas mensais de Caixa/Banco da igreja e gera os consolidados
Banco_<ano>.xlsx e Caixa_<ano>.xlsx no layout do sistema de importacao.

Uso:
  python processar.py --uploads /mnt/user-data/uploads --out /mnt/user-data/outputs --ano 2025
"""
import argparse, datetime, json, os, glob, sys
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

HEADERS = ['Data','Tipo','Método','Descrição','Categoria','Sub-Categoria','Código Sub-Categoria','Valor']

def load_classif(path):
    with open(path, encoding='utf-8') as f:
        raw = json.load(f)
    # keys are 3-digit strings -> int for lookup, keep formatted version too
    return {int(k): (k, v[0], v[1], v[2]) for k, v in raw.items()}

def norm_code(cell):
    """Converte a celula de codigo (numero ou texto RESGATE/APLICACAO) em int."""
    if cell is None:
        return None
    if isinstance(cell, (int, float)) and float(cell) == int(cell):
        return int(cell)
    if isinstance(cell, str):
        s = cell.strip().upper()
        if s == 'RESGATE' or s.startswith('RESGATE APLIC'):
            return 902
        if s == 'APLICAÇÃO' or s == 'APLICACAO':
            return 903
    return None

def detect_transfer_from_hist(hist):
    if not isinstance(hist, str):
        return None
    h = hist.strip().upper()
    if h == 'APLICAÇÃO' or h == 'APLICACAO' or h.startswith('APLICAÇÃO ') or h.startswith('APLICACAO '):
        return 903
    if h == 'RESGATE' or h.startswith('RESGATE APLIC'):
        return 902
    if 'TRANSFER' in h and 'CRESOL' in h:
        return 901
    return None

# Regras por DESCRIÇÃO: forcam o codigo independentemente do codigo numerico da origem.
# Avaliadas em ordem; a primeira que casar vence.
def _norm(s):
    import unicodedata
    s = unicodedata.normalize('NFKD', s).encode('ascii', 'ignore').decode().lower()
    return s

DESC_RULES = [
    (['ramon e ana'], 205),
    (['pense laranja'], 303),
    (['missionarios cobim'], 204),
    (['agencia missionaria'], 207),
    (['projeto missionario natal'], 204),
    (['projeto missionario picarras'], 204),
    (['projeto missionario itapema'], 204),
    (['projeto missionario itajai'], 204),
]

def code_from_description(hist):
    if not isinstance(hist, str):
        return None
    h = _norm(hist)
    for needles, code in DESC_RULES:
        if all(n in h for n in needles):
            return code
    return None


def find_header_row(ws):
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if row and row[0] == 'Dia/Mes':
            return i
    return 6

def parse_sheet(ws, kind):
    """kind: 'CX' ou 'BC'. Retorna lista de dicts."""
    if kind == 'CX':
        hi, ci, cr, db = 1, 2, 4, 5
    else:
        hi, ci, cr, db = 2, 3, 5, 6
    out = []
    start = find_header_row(ws) + 1
    for row in ws.iter_rows(min_row=start, values_only=True):
        if not row or not isinstance(row[0], datetime.datetime):
            continue
        hist, code, cred, deb = row[hi], row[ci], row[cr], row[db]
        if (cred is None or cred == 0) and (deb is None or deb == 0):
            continue
        nc = norm_code(code)
        if nc is None:
            nc = detect_transfer_from_hist(hist)
        # Regra por descricao tem prioridade (corrige codigos antigos/errados)
        desc_code = code_from_description(hist)
        if desc_code is not None:
            nc = desc_code
        out.append({
            'date': row[0].date(),
            'hist': hist.strip() if isinstance(hist, str) else (hist or ''),
            'code': nc, 'cred': cred, 'deb': deb,
        })
    return out

def collect(uploads, kind_label):
    """kind_label: 'Caixa' ou 'Banco'."""
    pref = 'Cx' if kind_label == 'Caixa' else 'Bc'
    k = 'CX' if kind_label == 'Caixa' else 'BC'
    rows = []
    for fp in sorted(glob.glob(os.path.join(uploads, '*.xlsx'))):
        try:
            wb = openpyxl.load_workbook(fp, data_only=True)
        except Exception:
            continue
        for sn in wb.sheetnames:
            if sn.lower().startswith(pref.lower()):
                rows.append((fp, sn, parse_sheet(wb[sn], k)))
    return rows

def build(uploads, kind_label, metodo, out_path, classif):
    sheets = collect(uploads, kind_label)
    wb = Workbook(); ws = wb.active; ws.title = kind_label

    HDR_FILL = PatternFill('solid', fgColor='1F3864')
    HDR_FONT = Font(name='Arial', bold=True, color='FFFFFF', size=11)
    CELL_FONT = Font(name='Arial', size=10)
    YELLOW = PatternFill('solid', fgColor='FFFF00')
    thin = Side(style='thin', color='D9D9D9')
    BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

    for j, h in enumerate(HEADERS, 1):
        c = ws.cell(1, j, h)
        c.fill = HDR_FILL; c.font = HDR_FONT
        c.alignment = Alignment(horizontal='center', vertical='center'); c.border = BORDER

    r = 2
    yellow_count = 0
    unknown = {}
    total = 0.0
    for _fp, _sn, data in sheets:
        for d in data:
            code = d['code']
            info = classif.get(code)
            if info:
                fcode, tipo, cat, sub = info
            else:
                tipo = 'Receita' if (d['cred'] not in (None, 0)) else 'Despesa'
                cat = ''; sub = None; fcode = None
                if code is not None:
                    unknown[code] = unknown.get(code, 0) + 1
            if d['cred'] not in (None, 0):
                val = abs(d['cred'])
            elif d['deb'] not in (None, 0):
                val = -abs(d['deb'])
            else:
                val = 0
            # Neutra: sinal segue credito(+)/debito(-) ja aplicado acima
            total += val

            ws.cell(r, 1, d['date']).number_format = 'DD/MM/YYYY'
            ws.cell(r, 2, tipo)
            ws.cell(r, 3, metodo)
            ws.cell(r, 4, d['hist'])
            ws.cell(r, 5, cat)
            sub_cell = ws.cell(r, 6, sub)
            code_cell = ws.cell(r, 7, fcode)
            if fcode is not None:
                code_cell.number_format = '@'
            vcell = ws.cell(r, 8, val)
            vcell.number_format = '#,##0.00;[Red](#,##0.00)'
            if info is None:
                sub_cell.fill = YELLOW; code_cell.fill = YELLOW
                yellow_count += 1
            for j in range(1, 9):
                cc = ws.cell(r, j)
                cc.font = CELL_FONT; cc.border = BORDER
                cc.alignment = Alignment(vertical='center',
                                         horizontal=('left' if j in (4, 5, 6) else 'center'))
            r += 1

    widths = {'A': 12, 'B': 10, 'C': 10, 'D': 52, 'E': 16, 'F': 40, 'G': 18, 'H': 14}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A2'
    if r > 2:
        ws.auto_filter.ref = f'A1:H{r-1}'
    wb.save(out_path)
    return {'arquivo': out_path, 'linhas': r - 2, 'amarelas': yellow_count,
            'codigos_desconhecidos': unknown, 'soma_valor': round(total, 2)}

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--uploads', default='/mnt/user-data/uploads')
    ap.add_argument('--out', default='/mnt/user-data/outputs')
    ap.add_argument('--ano', default='2025')
    ap.add_argument('--classif', default=None)
    args = ap.parse_args()

    here = os.path.dirname(os.path.abspath(__file__))
    classif_path = args.classif or os.path.join(here, '..', 'references', 'classificacao.json')
    classif = load_classif(classif_path)

    os.makedirs(args.out, exist_ok=True)
    rb = build(args.uploads, 'Banco', 'Banco',
               os.path.join(args.out, f'Banco_{args.ano}.xlsx'), classif)
    rc = build(args.uploads, 'Caixa', 'Caixa',
               os.path.join(args.out, f'Caixa_{args.ano}.xlsx'), classif)
    print(json.dumps({'Banco': rb, 'Caixa': rc}, ensure_ascii=False, indent=2))

if __name__ == '__main__':
    main()
