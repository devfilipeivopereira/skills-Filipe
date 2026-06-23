from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


WORD_RE = re.compile(r"\b[\w'-]+\b", re.UNICODE)


def load_config(config_path: Path) -> dict[str, Any]:
    return json.loads(config_path.read_text(encoding="utf-8"))


def count_words(text: str) -> int:
    return len(WORD_RE.findall(text))


def parse_blocks(text: str) -> tuple[list[dict[str, str]], str]:
    blocks: list[dict[str, str]] = []
    title = ""
    paragraph_lines: list[str] = []

    def flush_paragraph() -> None:
        if paragraph_lines:
            content = " ".join(line.strip() for line in paragraph_lines).strip()
            if content:
                blocks.append({"type": "paragraph", "text": content})
            paragraph_lines.clear()

    for raw_line in text.splitlines():
        stripped = raw_line.strip()

        if not stripped:
            flush_paragraph()
            continue

        if stripped.startswith("# "):
            flush_paragraph()
            text_value = stripped[2:].strip()
            if not title:
                title = text_value
            blocks.append({"type": "title", "text": text_value})
            continue

        if stripped.startswith("## "):
            flush_paragraph()
            blocks.append({"type": "heading1", "text": stripped[3:].strip()})
            continue

        if stripped.startswith("### "):
            flush_paragraph()
            blocks.append({"type": "heading2", "text": stripped[4:].strip()})
            continue

        if stripped.startswith("- "):
            flush_paragraph()
            blocks.append({"type": "bullet", "text": stripped[2:].strip()})
            continue

        paragraph_lines.append(stripped)

    flush_paragraph()
    if not title:
        title = "Sermon Manuscript"
    return blocks, title


def ensure_style(document: Document, style_name: str, base_name: str | None = None) -> None:
    if style_name in document.styles:
        return

    base_style = document.styles[base_name] if base_name else None
    style = document.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
    if base_style is not None:
        style.base_style = base_style


def set_run_font(run, font_name: str, size_pt: int, bold: bool = False) -> None:
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(size_pt)
    run.font.bold = bold


def configure_document(document: Document, config: dict[str, Any]) -> None:
    section = document.sections[0]
    margins = config["document"]["margins_pt"]
    section.top_margin = Pt(margins["top"])
    section.bottom_margin = Pt(margins["bottom"])
    section.left_margin = Pt(margins["left"])
    section.right_margin = Pt(margins["right"])

    fonts = config["document"]["fonts"]
    sizes = config["document"]["font_sizes_pt"]

    normal = document.styles["Normal"]
    normal.font.name = fonts["body"]
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), fonts["body"])
    normal.font.size = Pt(sizes["body"])
    normal.paragraph_format.space_after = Pt(10)
    normal.paragraph_format.line_spacing = 1.15

    title_style = document.styles["Title"]
    title_style.font.name = fonts["title"]
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), fonts["title"])
    title_style.font.size = Pt(sizes["title"])
    title_style.font.bold = True
    title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_style.paragraph_format.space_after = Pt(18)

    heading1 = document.styles["Heading 1"]
    heading1.font.name = fonts["heading"]
    heading1._element.rPr.rFonts.set(qn("w:eastAsia"), fonts["heading"])
    heading1.font.size = Pt(sizes["heading1"])
    heading1.font.bold = True
    heading1.paragraph_format.space_before = Pt(16)
    heading1.paragraph_format.space_after = Pt(8)

    heading2 = document.styles["Heading 2"]
    heading2.font.name = fonts["heading"]
    heading2._element.rPr.rFonts.set(qn("w:eastAsia"), fonts["heading"])
    heading2.font.size = Pt(sizes["heading2"])
    heading2.font.bold = True
    heading2.paragraph_format.space_before = Pt(12)
    heading2.paragraph_format.space_after = Pt(6)

    ensure_style(document, "Sermon Bullet", "List Bullet")
    bullet = document.styles["Sermon Bullet"]
    bullet.font.name = fonts["body"]
    bullet._element.rPr.rFonts.set(qn("w:eastAsia"), fonts["body"])
    bullet.font.size = Pt(sizes["body"])
    bullet.paragraph_format.space_after = Pt(4)


def add_header_footer(document: Document, title: str, config: dict[str, Any]) -> None:
    section = document.sections[0]
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run(title)
    set_run_font(header_run, config["document"]["fonts"]["body"], 9)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run(config["document"]["footer_text"])
    set_run_font(footer_run, config["document"]["fonts"]["body"], 9)


def render_docx_bytes(text: str, config: dict[str, Any]) -> tuple[bytes, int]:
    word_count = count_words(text)
    blocks, title = parse_blocks(text)

    document = Document()
    configure_document(document, config)
    add_header_footer(document, title, config)

    for block in blocks:
        block_type = block["type"]
        block_text = block["text"]
        if block_type == "title":
            document.add_paragraph(block_text, style="Title")
        elif block_type == "heading1":
            document.add_paragraph(block_text, style="Heading 1")
        elif block_type == "heading2":
            document.add_paragraph(block_text, style="Heading 2")
        elif block_type == "bullet":
            document.add_paragraph(block_text, style="Sermon Bullet")
        else:
            paragraph = document.add_paragraph(block_text, style="Normal")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    document.add_section(WD_SECTION.CONTINUOUS)

    from io import BytesIO

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue(), word_count
