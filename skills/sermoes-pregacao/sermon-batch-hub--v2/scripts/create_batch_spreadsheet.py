#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import json
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET


WORD_NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
WORD_RE = re.compile(r"\b[\w'-]+\b", re.UNICODE)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Create the sermon-batch summary spreadsheet for a centralized batch folder. "
            "Use --rows-json when the orchestrator has preacher metadata such as title, thesis, and distinctive resource."
        )
    )
    parser.add_argument("--reference", required=True, help="Biblical reference folder name, e.g. João_9")
    parser.add_argument(
        "--rows-json",
        default="",
        help="Optional path to a JSON file containing row metadata for the batch",
    )
    parser.add_argument(
        "--docs-root",
        default=r"C:\Users\filip\.codex\skills\DOCS_Sermons",
        help="Centralized DOCS_Sermons root",
    )
    parser.add_argument(
        "--epub-root",
        default=r"C:\Users\filip\.codex\skills\EPUB_SERMONS",
        help="Centralized EPUB_SERMONS root",
    )
    parser.add_argument(
        "--output-base-name",
        default="batch_summary",
        help="ASCII-safe base filename for the generated CSV and XLSX",
    )
    return parser.parse_args()


def sanitize_filename_part(value: str) -> str:
    cleaned = re.sub(r"[^\w]+", "_", value.strip(), flags=re.UNICODE)
    cleaned = re.sub(r"_+", "_", cleaned, flags=re.UNICODE).strip("_")
    return cleaned or "sem_nome"


def count_docx_words(docx_path: Path) -> int:
    with zipfile.ZipFile(docx_path) as zf:
        xml_bytes = zf.read("word/document.xml")
    root = ET.fromstring(xml_bytes)
    tokens: list[str] = []
    for node in root.findall(".//w:t", WORD_NS):
        if node.text:
            tokens.append(node.text)
    return len(WORD_RE.findall(" ".join(tokens)))


def load_rows(rows_json: str) -> list[dict[str, object]]:
    if not rows_json:
        return []
    payload = json.loads(Path(rows_json).read_text(encoding="utf-8"))
    if not isinstance(payload, list):
        raise ValueError("--rows-json must point to a JSON list")
    normalized: list[dict[str, object]] = []
    for row in payload:
        if not isinstance(row, dict):
            raise ValueError("Each row in --rows-json must be an object")
        normalized.append(row)
    return normalized


def infer_rows_from_outputs(reference: str, docs_dir: Path, epub_dir: Path) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for docx_path in sorted(docs_dir.glob("*.docx")):
        stem = docx_path.stem
        preacher = stem.removeprefix(f"{reference}_").replace("_", " ")
        epub_match = epub_dir / f"{stem}.epub"
        rows.append(
            {
                "Pregador": preacher,
                "Skill": "",
                "Titulo": "",
                "Tese_Central": "",
                "Recurso_Distintivo": "",
                "Min_Palavras": "",
                "DOCX": str(docx_path),
                "EPUB": str(epub_match) if epub_match.exists() else "",
            }
        )
    return rows


def enrich_rows(
    reference: str,
    rows: list[dict[str, object]],
    docs_dir: Path,
    epub_dir: Path,
) -> list[dict[str, object]]:
    enriched: list[dict[str, object]] = []
    for row in rows:
        preacher = str(row.get("Pregador") or row.get("preacher") or "").strip()
        skill = str(row.get("Skill") or row.get("skill") or "").strip()
        title = str(row.get("Titulo") or row.get("title") or "").strip()
        thesis = str(row.get("Tese_Central") or row.get("central_thesis") or row.get("refrain") or "").strip()
        resource = str(row.get("Recurso_Distintivo") or row.get("distinct_resource") or row.get("resource") or "").strip()
        min_words = row.get("Min_Palavras") or row.get("min_words") or ""

        docx_path_value = str(row.get("DOCX") or row.get("docx") or "").strip()
        if docx_path_value:
            docx_path = Path(docx_path_value)
        else:
            docx_name = f"{reference}_{sanitize_filename_part(preacher)}.docx"
            docx_path = docs_dir / docx_name

        epub_path_value = str(row.get("EPUB") or row.get("epub") or "").strip()
        if epub_path_value:
            epub_path = Path(epub_path_value)
        else:
            expected = epub_dir / f"{docx_path.stem}.epub"
            if expected.exists():
                epub_path = expected
            else:
                candidates = sorted(epub_dir.glob(f"*{sanitize_filename_part(preacher)}*.epub"))
                epub_path = candidates[0] if candidates else Path()

        total_words = count_docx_words(docx_path) if docx_path.exists() else ""
        enriched.append(
            {
                "Pregador": preacher,
                "Skill": skill,
                "Titulo": title,
                "Tese_Central": thesis,
                "Recurso_Distintivo": resource,
                "Min_Palavras": min_words,
                "Total_Palavras": total_words,
                "DOCX": str(docx_path) if docx_path else "",
                "EPUB": str(epub_path) if str(epub_path) not in {".", ""} else "",
            }
        )
    return enriched


def write_csv(csv_path: Path, rows: list[dict[str, object]]) -> None:
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = list(rows[0].keys()) if rows else [
        "Pregador",
        "Skill",
        "Titulo",
        "Tese_Central",
        "Recurso_Distintivo",
        "Min_Palavras",
        "Total_Palavras",
        "DOCX",
        "EPUB",
    ]
    with csv_path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, delimiter=";")
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx(xlsx_path: Path, rows: list[dict[str, object]]) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except Exception:
        return False

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    headers = list(rows[0].keys()) if rows else [
        "Pregador",
        "Skill",
        "Titulo",
        "Tese_Central",
        "Recurso_Distintivo",
        "Min_Palavras",
        "Total_Palavras",
        "DOCX",
        "EPUB",
    ]
    wb = Workbook()
    ws = wb.active
    ws.title = "Batch Summary"
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header, "") for header in headers])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = {
        "A": 22,
        "B": 36,
        "C": 36,
        "D": 42,
        "E": 50,
        "F": 14,
        "G": 16,
        "H": 85,
        "I": 85,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wb.save(xlsx_path)
    return True


def main() -> int:
    args = parse_args()
    docs_dir = Path(args.docs_root) / args.reference
    epub_dir = Path(args.epub_root) / args.reference
    rows = load_rows(args.rows_json)
    if not rows:
        rows = infer_rows_from_outputs(args.reference, docs_dir, epub_dir)
    rows = enrich_rows(args.reference, rows, docs_dir, epub_dir)

    base_name = sanitize_filename_part(args.output_base_name)
    csv_path = docs_dir / f"{base_name}.csv"
    xlsx_path = docs_dir / f"{base_name}.xlsx"

    write_csv(csv_path, rows)
    xlsx_written = write_xlsx(xlsx_path, rows)

    print(f"CSV={csv_path}")
    if xlsx_written:
        print(f"XLSX={xlsx_path}")
    else:
        print("XLSX=not_written")
    print(f"ROWS={len(rows)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
