from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


ABOUT_AUTHOR_TEXT = (
    "Filipe Ivo Pereira é pastor, escritor, mentor e docente, dedicado a ajudar pessoas a conectarem fé, "
    "sabedoria bíblica e vida prática. Natural de Florianópolis, Santa Catarina, também atua nas áreas de "
    "finanças e estratégia, reunindo em sua trajetória experiência ministerial, sólida formação acadêmica e "
    "sensibilidade para os desafios do mundo contemporâneo.\n\n"
    "Mestre em Administração pela ESAG/UDESC e com MBA em Gestão de Investimentos, possui formação em Teologia, "
    "Ciências Contábeis e Administração, além de especializações em campos como ciência de dados, inteligência "
    "artificial, engenharia de dados, business intelligence e engenharia de software. Também é graduado em "
    "Liderança Avançada pelo Haggai Brasil.\n\n"
    "Desde 2009, serve como pastor na Igreja da Família Cristã, em São José (SC), sendo ordenado pela Convenção "
    "Brasileira das Igrejas Irmãos Menonitas (COBIM). Atua ainda como professor na Faculdade Fidelis, no Instituto "
    "Haggai do Brasil e em cursos voltados à liderança, gestão e espiritualidade cristã.\n\n"
    "Autor de dezenas de livros, também escreve no blog Atitude Invest, onde compartilha reflexões e conteúdos "
    "sobre educação financeira e investimentos.\n\n"
    "Sua paixão é ver o Evangelho de Cristo transformando não apenas indivíduos, mas famílias e comunidades "
    "inteiras. Casado com Renata da Silva Cardoso Pereira e pai de Daniel e Amanda, dedica sua vida a servir com "
    "equilíbrio, propósito e fé.\n\n"
    "Para conhecer mais materiais de estudo, ebooks e recursos pastorais, visite: "
    "[www.filipeivopereira.com](https://www.filipeivopereira.com)"
)
COPYRIGHT_LINES = [
    "{title}",
    "Filipe Ivo Pereira",
    "www.filipeivopereira.com",
    "Copyright © 2026 | Todos direitos reservados",
]
BLACK = RGBColor(0, 0, 0)
BODY_LINE_SPACING = 1.25
BODY_FIRST_LINE_INDENT = Mm(8)
BIBLE_QUOTE_FONT = "Times New Roman"


@dataclass
class Section:
    title: str
    paragraphs: list[str]


_bookmark_id = 1


def parse_markdown_sections(markdown_text: str, default_title: str) -> tuple[str, list[Section]]:
    title = default_title
    sections: list[Section] = []
    current_title: str | None = None
    current_lines: list[str] = []

    def flush() -> None:
        nonlocal current_title, current_lines
        if current_title:
            body = "\n".join(current_lines).strip()
            paragraphs = [piece.strip() for piece in re.split(r"\n\s*\n", body) if piece.strip()]
            sections.append(Section(current_title.strip(), paragraphs))
        current_title = None
        current_lines = []

    normalized = markdown_text.replace("\r\n", "\n").replace("\r", "\n")
    for line in normalized.split("\n"):
        if line.startswith("# "):
            title = line[2:].strip() or title
            continue
        if line.startswith("## "):
            flush()
            current_title = line[3:].strip()
            continue
        if current_title is not None:
            current_lines.append(line)
    flush()
    return title, sections


def render_markdown_docx(markdown_text: str, output_path: Path, *, fallback_title: str) -> Path:
    global _bookmark_id
    _bookmark_id = 1
    title, sections = parse_markdown_sections(markdown_text, fallback_title)
    if not sections:
        raise ValueError("O markdown precisa conter headings `##` para as seções do livro.")

    preface = _pick_section(sections, "prefácio") or Section(
        "Prefácio",
        [
            f"{title} nasce de um sermão e ganha aqui forma de livro. O objetivo não é substituir a mensagem original, "
            "mas torná-la mais legível, mais contemplativa e mais útil para leitura continuada.",
        ],
    )
    intro = _pick_section(sections, "introdução") or Section(
        "Introdução",
        [
            "Leia com calma, destaque frases importantes e transforme convicção em resposta prática. "
            "Este livro foi organizado para conduzir você pela mensagem com clareza, reverência e aplicação pastoral.",
        ],
    )
    about = Section(
        "Sobre o Autor",
        [paragraph.strip() for paragraph in ABOUT_AUTHOR_TEXT.split("\n\n") if paragraph.strip()],
    )

    special_titles = {preface.title.lower(), intro.title.lower(), about.title.lower()}
    body_sections = [s for s in sections if s.title.lower() not in special_titles]

    doc = Document()
    _render_title_page(doc, title)
    _render_copyright_page(doc, title)
    _render_toc(doc)
    _render_section(doc, preface, drop_cap=False)
    doc.add_page_break()
    _render_section(doc, intro, drop_cap=False)

    for index, section in enumerate(body_sections):
        doc.add_page_break()
        _render_section(doc, section, drop_cap=True)
        if index < len(body_sections) - 1:
            doc.add_page_break()
            _render_quote_page(doc, section)

    doc.add_page_break()
    _render_section(doc, about, drop_cap=False)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path


def _pick_section(sections: list[Section], title: str) -> Section | None:
    for section in sections:
        if section.title.strip().lower() == title:
            return section
    return None


def _render_title_page(doc: Document, title: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(220)
    run = p.add_run(title)
    run.font.size = Pt(28)
    run.bold = True
    run.font.color.rgb = BLACK
    doc.add_page_break()


def _render_copyright_page(doc: Document, title: str) -> None:
    for idx, line in enumerate(COPYRIGHT_LINES):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if idx == 0:
            p.paragraph_format.space_before = Pt(180)
        if line == "www.filipeivopereira.com":
            _add_external_link(p, "https://www.filipeivopereira.com", line)
            continue
        run = p.add_run(line.format(title=title))
        if idx == 0:
            run.bold = True
        run.font.color.rgb = BLACK
    doc.add_page_break()


def _render_toc(doc: Document) -> None:
    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.space_before = Pt(16)
    heading.paragraph_format.space_after = Pt(12)
    heading_run = heading.add_run("SUMÁRIO")
    heading_run.bold = True
    heading_run.font.size = Pt(16)
    heading_run.font.color.rgb = BLACK

    toc_paragraph = doc.add_paragraph()
    _insert_toc_field(toc_paragraph)
    doc.add_page_break()


_EFEITO_RE = re.compile(r"<!--\s*EFEITO:\s*(.*?)\s*-->", re.IGNORECASE | re.DOTALL)
_EFEITO_FALLBACK = "A graça fala mais alto do que qualquer queda; e isso muda o que você faz com o que ainda tem."


def _extract_effect_phrase(section: Section) -> tuple[list[str], str]:
    """Return (clean_paragraphs_without_efeito_comment, effect_phrase_text).
    The EFEITO comment is written by the AI agent at the end of each chapter as:
      <!-- EFEITO: frase aqui -->
    The renderer reads it, strips it from body content, and uses it on the quote page.
    Falls back to _EFEITO_FALLBACK if the AI did not include the marker.
    """
    clean: list[str] = []
    phrase: str | None = None
    for p in section.paragraphs:
        m = _EFEITO_RE.search(p)
        if m:
            candidate = m.group(1).strip()
            if candidate:
                phrase = candidate
        else:
            clean.append(p)
    return clean, phrase or _EFEITO_FALLBACK


def _render_quote_page(doc: Document, section: Section) -> None:
    _, sentence = _extract_effect_phrase(section)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(220)
    run = p.add_run(f"\"{sentence}\"")
    run.italic = True
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = BLACK


def _render_section(doc: Document, section: Section, *, drop_cap: bool) -> None:
    clean_paragraphs, _ = _extract_effect_phrase(section)
    heading = doc.add_heading(section.title, level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in heading.runs:
        run.font.color.rgb = BLACK
    _add_bookmark(heading, _slug(section.title))
    first_body = drop_cap
    for paragraph in clean_paragraphs:
        cleaned = paragraph.strip()
        if not cleaned:
            continue
        if cleaned.startswith("### "):
            _render_subheading(doc, cleaned[4:].strip())
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = BODY_LINE_SPACING
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.first_line_indent = BODY_FIRST_LINE_INDENT
        if first_body:
            first_body = False
            p.paragraph_format.first_line_indent = 0
            first = p.add_run(cleaned[0])
            first.font.size = Pt(34)
            first.font.color.rgb = BLACK
            first.bold = True
            _append_mixed_content(p, cleaned[1:])
        else:
            _append_mixed_content(p, cleaned)


def _render_subheading(doc: Document, text: str) -> None:
    p = doc.add_heading(text, level=2)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 2.0
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(12)
    for run in p.runs:
        run.bold = True
        run.font.size = Pt(14)
        run.font.color.rgb = BLACK



def _next_bookmark_id() -> str:
    global _bookmark_id
    value = str(_bookmark_id)
    _bookmark_id += 1
    return value


def _slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_") or "secao"


def _add_bookmark(paragraph, name: str) -> None:
    bookmark_id = _next_bookmark_id()
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), bookmark_id)
    start.set(qn("w:name"), name)
    paragraph._p.append(start)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), bookmark_id)
    paragraph._p.append(end)


def _insert_toc_field(paragraph) -> None:
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")

    instr_run = OxmlElement("w:r")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = 'TOC \\o "1-3" \\h \\z \\u'
    instr_run.append(instr_text)

    fld_separate = OxmlElement("w:fldChar")
    fld_separate.set(qn("w:fldCharType"), "separate")

    placeholder_run = OxmlElement("w:r")
    placeholder_text = OxmlElement("w:t")
    placeholder_text.text = "Atualize os campos no Word para montar o sumario automatico."
    placeholder_run.append(placeholder_text)

    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    paragraph._p.append(fld_begin)
    paragraph._p.append(instr_run)
    paragraph._p.append(fld_separate)
    paragraph._p.append(placeholder_run)
    paragraph._p.append(fld_end)


def _add_internal_link(paragraph, anchor: str, text: str) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    hyperlink.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    bold = OxmlElement("w:b")
    rpr.append(bold)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _append_mixed_content(paragraph, text: str) -> None:
    link_pattern = re.compile(r"\[([^\]]+)\]\((https?://[^)]+)\)")
    pos = 0
    for match in link_pattern.finditer(text):
        if match.start() > pos:
            _append_emphasis_runs(paragraph, text[pos:match.start()])
        label, url = match.group(1), match.group(2)
        _add_external_link(paragraph, url, label)
        pos = match.end()
    if pos < len(text):
        _append_emphasis_runs(paragraph, text[pos:])


def _append_emphasis_runs(paragraph, text: str) -> None:
    emphasis_pattern = re.compile(r"\*([^*\n]+)\*")
    pos = 0
    for match in emphasis_pattern.finditer(text):
        if match.start() > pos:
            _add_text_run(paragraph, text[pos:match.start()])
        _add_text_run(paragraph, match.group(1), italic=True, font_name=BIBLE_QUOTE_FONT)
        pos = match.end()
    if pos < len(text):
        _add_text_run(paragraph, text[pos:])


def _add_text_run(paragraph, text: str, *, italic: bool = False, font_name: str | None = None):
    run = paragraph.add_run(text)
    run.font.color.rgb = BLACK
    if italic:
        run.italic = True
    if font_name:
        run.font.name = font_name
        r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
        r_fonts.set(qn("w:ascii"), font_name)
        r_fonts.set(qn("w:hAnsi"), font_name)
        r_fonts.set(qn("w:cs"), font_name)
    return run


def _add_external_link(paragraph, url: str, text: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    run.append(rpr)

    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
